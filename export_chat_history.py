import json
import re
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

transcript_path = r"C:\Users\rchuk\.gemini\antigravity\brain\797e1d1a-07d3-4d1b-a401-45de60e4324a\.system_generated\logs\transcript_full.jsonl"
lab_dir = r"c:\Users\rchuk\Downloads\09_tic_tac_toe\Lab-4"
os.makedirs(lab_dir, exist_ok=True)

messages = []

with open(transcript_path, "r", encoding="utf-8") as f:
    for line in f:
        try:
            data = json.loads(line)
        except Exception:
            continue
        
        stype = data.get("type")
        source = data.get("source")
        content = data.get("content", "")
        
        # User input
        if stype == "USER_INPUT" or source == "USER_EXPLICIT":
            # Clean out XML tags if any
            clean_content = re.sub(r"<USER_REQUEST>(.*?)</USER_REQUEST>", r"\1", content, flags=re.DOTALL)
            clean_content = re.sub(r"<ADDITIONAL_METADATA>.*?</ADDITIONAL_METADATA>", "", clean_content, flags=re.DOTALL)
            clean_content = re.sub(r"<USER_SETTINGS_CHANGE>.*?</USER_SETTINGS_CHANGE>", "", clean_content, flags=re.DOTALL)
            clean_content = re.sub(r"==Start of PDF==.*?==End of PDF==", "[Uploaded Lab 4 PDF: VibeCoding Instructions]", clean_content, flags=re.DOTALL)
            clean_content = clean_content.strip()
            if clean_content:
                messages.append({"role": "User", "text": clean_content})
        
        # Assistant response
        elif stype == "PLANNER_RESPONSE" or source == "MODEL":
            if content and not data.get("tool_calls"):
                clean_content = content.strip()
                if clean_content:
                    messages.append({"role": "Antigravity AI (Pair Programmer)", "text": clean_content})

print(f"Total extracted conversation turns: {len(messages)}")

# Generate Markdown
md_path = os.path.join(lab_dir, "PES1UG24CS135_Lab4_Chat_History.md")
with open(md_path, "w", encoding="utf-8") as f:
    f.write("# Lab 4: VibeCoding — Complete Chat History\n\n")
    f.write("**Student Name:** Rohan Chukkapalli  \n")
    f.write("**PRN / SRN:** PES1UG24CS135  \n")
    f.write("**Repository:** https://github.com/SETAPESU26/09_tic_tac_toe  \n")
    f.write("**Personal Repo:** https://github.com/VortexVector/09_tic_tac_toe  \n\n")
    f.write("---\n\n")
    for msg in messages:
        f.write(f"### {msg['role']}\n\n{msg['text']}\n\n---\n\n")

print(f"Saved MD to {md_path}")

# Generate DOCX
doc = Document()
title = doc.add_heading("Lab 4: VibeCoding — Chat History", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Student Name: Rohan Chukkapalli | PRN: PES1UG24CS135\n").bold = True
meta.add_run("Course Assignment: Lab 4 VibeCoding\n")
meta.add_run("GitHub: https://github.com/VortexVector/09_tic_tac_toe\n")

for msg in messages:
    h = doc.add_heading(msg["role"], level=2)
    p = doc.add_paragraph(msg["text"])
    doc.add_paragraph("-" * 40)

docx_path = os.path.join(lab_dir, "PES1UG24CS135_Lab4_Chat_History.docx")
doc.save(docx_path)
print(f"Saved DOCX to {docx_path}")

# Generate PDF
pdf_path = os.path.join(lab_dir, "PES1UG24CS135_Lab4_Chat_History.pdf")
doc_pdf = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
styles = getSampleStyleSheet()

role_user_style = ParagraphStyle(
    "RoleUser",
    parent=styles["Heading2"],
    fontSize=11,
    leading=14,
    textColor=colors.HexColor("#2B6CB0"),
    spaceBefore=8,
    spaceAfter=4
)
role_ai_style = ParagraphStyle(
    "RoleAI",
    parent=styles["Heading2"],
    fontSize=11,
    leading=14,
    textColor=colors.HexColor("#2C5282"),
    spaceBefore=8,
    spaceAfter=4
)
body_style = ParagraphStyle(
    "ChatBody",
    parent=styles["Normal"],
    fontSize=8.5,
    leading=12,
    textColor=colors.HexColor("#1A202C"),
    spaceAfter=6
)

story = []
story.append(Paragraph("Lab 4: VibeCoding — Complete Chat History", styles["Title"]))
story.append(Paragraph("<b>Student:</b> Rohan Chukkapalli (PES1UG24CS135) &nbsp;|&nbsp; <b>Repo:</b> VortexVector/09_tic_tac_toe", styles["Normal"]))
story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#CBD5E0"), spaceAfter=10))

for msg in messages:
    style_header = role_user_style if msg["role"] == "User" else role_ai_style
    story.append(Paragraph(f"<b>{msg['role']}</b>", style_header))
    # Replace newlines with <br/> for ReportLab Paragraph
    text_formatted = msg["text"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
    story.append(Paragraph(text_formatted, body_style))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceAfter=8))

doc_pdf.build(story)
print(f"Saved PDF to {pdf_path}")
